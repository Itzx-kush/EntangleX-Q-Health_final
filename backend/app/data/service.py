from __future__ import annotations

import csv
import hashlib
import io
from threading import Lock
from uuid import uuid4

import numpy as np
import pandas as pd
from sqlalchemy import select

from ..api.schemas import DatasetUploadMetadata
from ..config import get_settings
from ..database import session_scope
from ..storage.entities import Dataset, Experiment
from ..storage.files import atomic_bytes, safe_path, sanitize_filename, verify
from ..storage.repository import require
from ..utils.errors import AppError
from ..utils.serialization import utcnow
from .catalog import benchmark_bytes, get_benchmark, list_benchmarks
from .metadata import suggest_metadata
from .quality import compatibility_report, inspection_report, quality_report

SUPPORTED_FORMATS = {".csv", ".tsv", ".xlsx", ".xls", ".parquet"}
_demo_registration_lock = Lock()
_dataset_delete_lock = Lock()


def _extension(filename: str) -> str:
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_FORMATS:
        supported = ", ".join(sorted(ext.lstrip(".").upper() for ext in SUPPORTED_FORMATS))
        raise AppError("extension_not_allowed", f"Unsupported dataset format. Use one of: {supported}.")
    return extension


def detect_file_format(content: bytes, filename: str) -> str:
    extension = _extension(filename)
    if b"\x00" in content and extension in {".csv", ".tsv"}:
        raise AppError("invalid_tabular_file", "Text tabular files cannot contain binary/NUL content.")
    if extension == ".xlsx" and not content.startswith(b"PK"):
        raise AppError("invalid_spreadsheet", "The .xlsx upload is not a valid ZIP-based Excel workbook.")
    if extension == ".xls" and not content.startswith(b"\xd0\xcf\x11\xe0"):
        raise AppError("invalid_spreadsheet", "The .xls upload is not a valid legacy Excel workbook.")
    if extension == ".parquet" and not (content.startswith(b"PAR1") and content.endswith(b"PAR1")):
        raise AppError("invalid_parquet", "The Parquet upload does not have a valid Parquet signature.")
    return extension


def _clean_header_names(columns, settings) -> list[str]:
    names = [str(column).lstrip("\ufeff").strip() for column in columns]
    if len(names) < 2 or len(names) > settings.max_columns:
        raise AppError("column_limit", "The dataset must contain a target and input features within the column limit.")
    if len(set(names)) != len(names) or any(not name or len(name) > 100 or any(ord(c) < 32 for c in name) for name in names):
        raise AppError("invalid_header", "Column names must be unique, trimmed, nonempty, and at most 100 characters.")
    return names


def _validate_headers(frame: pd.DataFrame, settings, header_names: list[str] | None = None) -> pd.DataFrame:
    if frame is None:
        raise AppError("column_limit", "The dataset must contain a target and input features within the column limit.")
    names = _clean_header_names(header_names if header_names is not None else frame.columns, settings)
    if len(names) != len(frame.columns):
        raise AppError("inconsistent_schema", "Rows and headers have inconsistent field counts.")
    frame = frame.copy()
    frame.columns = names
    if len(frame) < 10 or len(frame) > settings.max_rows:
        raise AppError("row_limit", "The dataset must have at least 10 rows and remain within the configured row limit.")
    return frame


def _spreadsheet_header(content: bytes, extension: str) -> list[str]:
    try:
        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                values = next(workbook.worksheets[0].iter_rows(values_only=True))
            finally:
                workbook.close()
            return list(values)
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            return workbook.sheet_by_index(0).row_values(0)
        finally:
            workbook.release_resources()
    except Exception as exc:
        raise AppError("invalid_spreadsheet", "The spreadsheet header could not be read safely.") from exc


def _normalize_object_columns(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    def normalize(value):
        # Keep complex values intact so the inspector can block them explicitly;
        # never stringify a nested value into a misleading categorical label.
        if isinstance(value, (list, dict, set, tuple, bytes, bytearray, complex)):
            return value
        try:
            return str(value) if pd.notna(value) else np.nan
        except (TypeError, ValueError):
            return value
    for column in frame.columns:
        if pd.api.types.is_object_dtype(frame[column]) or pd.api.types.is_string_dtype(frame[column]):
            frame[column] = frame[column].map(normalize).astype(object)
    return frame


def _parse_delimited(content: bytes, target: str | None, delimiter: str) -> pd.DataFrame:
    settings = get_settings()
    try:
        text = content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
        if not rows:
            raise StopIteration
        header = [cell.lstrip("\ufeff") for cell in rows[0]]
    except (UnicodeDecodeError, StopIteration, csv.Error) as exc:
        kind = "TSV" if delimiter == "\t" else "CSV"
        raise AppError("invalid_tabular_file", f"Upload a nonempty UTF-8 {kind} file with a consistent table schema.") from exc
    if len(header) < 2 or len(header) > settings.max_columns:
        raise AppError("column_limit", "The dataset must contain a target and input features within the column limit.")
    if len(set(header)) != len(header) or any(not h.strip() or h != h.strip() or len(h) > 100 or any(ord(c) < 32 for c in h) for h in header):
        raise AppError("invalid_header", "Column names must be unique, trimmed, nonempty, and at most 100 characters.")
    if any(len(row) != len(header) for row in rows[1:]):
        raise AppError("inconsistent_schema", "Rows and headers have inconsistent field counts.")
    dtype = {target: "string"} if target and target in header else None
    try:
        frame = pd.read_csv(io.StringIO(text), sep=delimiter, dtype=dtype, nrows=settings.max_rows + 1, on_bad_lines="error")
    except (ValueError, pd.errors.ParserError, pd.errors.EmptyDataError) as exc:
        raise AppError("invalid_tabular_file", "Tabular parsing failed; check quoting, delimiters, and the table schema.") from exc
    frame = _validate_headers(frame, settings)
    if list(frame.columns) != header:
        raise AppError("inconsistent_schema", "Rows and headers have inconsistent field counts.")

    # Preserve the existing, explicit diabetes readmission convenience mapping.
    readmission_source = "readmitted" in header and target and target.strip() in {"readmitted", "readmitted_30d"}
    if readmission_source:
        within_30 = frame["readmitted"].astype("string").str.strip().eq("<30")
        if target.strip() == "readmitted_30d" or (target not in header):
            frame[target] = within_30.astype(int)
            frame = frame.drop(columns=["readmitted", "encounter_id", "patient_nbr"], errors="ignore")
        else:
            frame["readmitted"] = np.where(within_30, "<30", "not_within_30d")
            frame = frame.drop(columns=["encounter_id", "patient_nbr"], errors="ignore")
    return _normalize_object_columns(frame)


def parse_tabular(content: bytes, filename: str, target: str | None = None, positive_label: str | None = None) -> pd.DataFrame:
    settings = get_settings()
    if len(content) > settings.upload_limit:
        raise AppError("upload_too_large", "The upload exceeds the configured size limit.", 413)
    extension = detect_file_format(content, filename)
    if extension == ".csv":
        return _parse_delimited(content, target, ",")
    if extension == ".tsv":
        return _parse_delimited(content, target, "\t")
    try:
        if extension in {".xlsx", ".xls"}:
            engine = "openpyxl" if extension == ".xlsx" else "xlrd"
            raw_header = _spreadsheet_header(content, extension)
            frame = pd.read_excel(io.BytesIO(content), sheet_name=0, engine=engine, nrows=settings.max_rows + 1)
            frame = _validate_headers(frame, settings, raw_header)
        else:
            frame = pd.read_parquet(io.BytesIO(content))
    except (ImportError, ModuleNotFoundError) as exc:
        raise AppError("reader_unavailable", f"The safe reader for {extension.lstrip('.').upper()} is not installed on the server.", 500) from exc
    except Exception as exc:
        code = "invalid_parquet" if extension == ".parquet" else "invalid_spreadsheet"
        label = "Parquet" if extension == ".parquet" else "spreadsheet"
        raise AppError(code, f"The {label} could not be parsed safely.") from exc
    if extension != ".xlsx" and extension != ".xls":
        frame = _validate_headers(frame, settings)
    if target and target in frame:
        frame[target] = frame[target].astype("string")
    return _normalize_object_columns(frame)


def parse_csv(content: bytes, target: str, positive_label: str | None = None) -> pd.DataFrame:
    return parse_tabular(content, "dataset.csv", target, positive_label)


def inspect_file(content: bytes, filename: str, target: str | None = None, positive_label: str | None = None) -> dict:
    extension = detect_file_format(content, filename)
    frame = parse_tabular(content, filename, target, positive_label)
    report = inspection_report(frame, target, positive_label)
    suggestions, metadata_warnings = suggest_metadata(
        filename,
        report,
        content_hash=hashlib.sha256(content).hexdigest(),
        selected_target=target,
        selected_positive=positive_label,
    )
    report["suggested_metadata"] = suggestions
    report["metadata_warnings"] = metadata_warnings
    report["file_format"] = extension.lstrip(".")
    report["filename"] = sanitize_filename(filename)
    report["file_size_bytes"] = len(content)
    return report


def _register_frame(content: bytes, filename: str, metadata: DatasetUploadMetadata, *, license_info: str | None = None, is_demo: bool = False, catalog_slug: str | None = None) -> Dataset:
    extension = detect_file_format(content, filename)
    frame = parse_tabular(content, filename, metadata.target, metadata.positive_label)
    quality = quality_report(frame, metadata.target, metadata.positive_label)
    compatibility = compatibility_report(frame, metadata.target, metadata.positive_label)
    inspection = inspection_report(frame, metadata.target, metadata.positive_label)
    catalog = get_benchmark(catalog_slug) if catalog_slug else None
    suggested_metadata, inferred_warnings = suggest_metadata(
        filename,
        inspection,
        content_hash=hashlib.sha256(content).hexdigest(),
        catalog=catalog,
        selected_target=metadata.target,
        selected_positive=metadata.positive_label,
    )
    if metadata.negative_label is not None and metadata.negative_label not in quality["target_classes"]:
        raise AppError("negative_label_unknown", "The configured negative label is not present in the target.")
    identity = str(uuid4())
    timestamp = utcnow()
    sha = hashlib.sha256(content).hexdigest()
    source_format = extension.lstrip(".")
    provenance = {
        **metadata.model_dump(exclude={"deidentified"}),
        "task": "binary_classification", "dataset_hash": sha, "hash_algorithm": "sha256",
        "hash_scope": "exact stored source bytes", "file_format": source_format,
        "file_extension": extension, "original_filename": sanitize_filename(filename),
        "uploaded_at": timestamp.isoformat(), "row_count": len(frame),
        "feature_count": len(frame.columns) - 1, "features": [c for c in frame if c != metadata.target],
        "numeric_features": quality["numeric_features"], "categorical_features": quality["categorical_features"],
        "class_distribution": quality["class_distribution"], "target_classes": quality["target_classes"],
        "negative_label": metadata.negative_label or next(c for c in quality["target_classes"] if c != metadata.positive_label),
        "license": license_info, "is_demo": is_demo, "catalog_slug": catalog_slug,
        "deidentification_asserted_by_uploader": not is_demo,
        "compatibility_findings": compatibility,
        "metadata_suggestions": metadata.metadata_suggestions or suggested_metadata,
        "metadata_warnings": list(dict.fromkeys(metadata.metadata_warnings + inferred_warnings)),
        "metadata_sources": metadata.metadata_sources or {key: value for key, value in suggested_metadata.items() if key.endswith("_source")},
        "metadata_confirmed": metadata.metadata_confirmed,
        "transformations_applied": ["Parsed into the internal pandas dataframe workflow; source bytes retained unchanged."],
        "preprocessing_configuration": "Stored per experiment; source dataset is immutable.",
    }
    path = safe_path("data/datasets", identity, extension)
    atomic_bytes(path, content)
    try:
        with session_scope() as session:
            record = Dataset(id=identity, name=metadata.name, filename=sanitize_filename(filename), sha256=sha, provenance=provenance, quality=quality, created_at=timestamp)
            session.add(record)
        return record
    except Exception:
        path.unlink(missing_ok=True)
        raise


def register_file(content: bytes, filename: str, metadata: DatasetUploadMetadata, *, license_info: str | None = None, is_demo: bool = False, catalog_slug: str | None = None) -> Dataset:
    return _register_frame(content, filename, metadata, license_info=license_info, is_demo=is_demo, catalog_slug=catalog_slug)


def register_csv(content: bytes, filename: str, metadata: DatasetUploadMetadata, *, license_info: str | None = None) -> Dataset:
    if not filename.lower().endswith(".csv"):
        raise AppError("extension_not_allowed", "The CSV registration helper accepts .csv files only.")
    return register_file(content, filename, metadata, license_info=license_info, is_demo=license_info is not None)


def load_frame(identity: str) -> tuple[Dataset, pd.DataFrame]:
    with session_scope() as session:
        record = require(session, Dataset, identity)
    extension = record.provenance.get("file_extension", ".csv")
    path = safe_path("data/datasets", identity, extension)
    verify(path, record.sha256)
    return record, parse_tabular(path.read_bytes(), record.filename, record.provenance["target"], record.provenance["positive_label"])


def list_library() -> list[dict]:
    return list_benchmarks()


def register_builtin(slug: str) -> Dataset:
    entry, content = benchmark_bytes(slug)
    metadata = DatasetUploadMetadata(
        name=entry["name"], domain=entry["domain"], source=entry["source"], source_url=entry["source_url"],
        version=entry["version"], target=entry["target"], positive_label=entry["positive_label"], deidentified=True,
    )
    with session_scope() as session:
        existing = session.scalar(select(Dataset).where(Dataset.sha256 == entry["dataset_hash"], Dataset.name == entry["name"]))
        if existing is not None:
            return existing
    return register_file(content, entry["filename"], metadata, license_info=entry["license"], is_demo=True, catalog_slug=entry["slug"])


def register_demo() -> Dataset:
    with _demo_registration_lock:
        return register_builtin("wdbc")


def delete_dataset(identity: str) -> None:
    with _dataset_delete_lock:
        _delete_dataset(identity)


def _delete_dataset(identity: str) -> None:
    with session_scope() as session:
        dataset = require(session, Dataset, identity)
        if session.scalar(select(Experiment.id).where(Experiment.dataset_id == identity).limit(1)):
            raise AppError("dataset_in_use", "This dataset is referenced by an experiment and is retained for reproducibility.", 409)
        expected = dataset.sha256
        extension = dataset.provenance.get("file_extension", ".csv")
    path = safe_path("data/datasets", identity, extension)
    if not path.is_file():
        raise AppError("integrity_error", "Stored dataset is missing and cannot be safely deleted.", 409)
    try:
        original = path.read_bytes()
    except OSError as exc:
        raise AppError("integrity_error", "Stored dataset could not be read and cannot be safely deleted.", 409) from exc
    if hashlib.sha256(original).hexdigest() != expected:
        raise AppError("integrity_error", "Stored dataset integrity has changed.", 409)
    try:
        path.unlink()
    except OSError as exc:
        raise AppError("storage_delete_failed", "Stored dataset could not be removed safely.", 409) from exc
    try:
        with session_scope() as session:
            dataset = require(session, Dataset, identity)
            if session.scalar(select(Experiment.id).where(Experiment.dataset_id == identity).limit(1)):
                raise AppError("dataset_in_use", "This dataset is referenced by an experiment and is retained for reproducibility.", 409)
            session.delete(dataset)
    except Exception:
        try:
            atomic_bytes(path, original)
        except OSError as exc:
            raise AppError("storage_consistency", "Dataset deletion failed and storage could not be restored.", 500) from exc
        raise
